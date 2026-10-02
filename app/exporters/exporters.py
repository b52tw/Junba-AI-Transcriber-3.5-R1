from __future__ import annotations
from pathlib import Path
from docx import Document
import html
import json
from app.core.models import TranscriptResult, Segment


def _time_srt(seconds: float) -> str:
    ms = max(0, int(round(seconds * 1000)))
    h, rem = divmod(ms, 3600000)
    m, rem = divmod(rem, 60000)
    s, ms = divmod(rem, 1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}'


def _time_vtt(seconds: float) -> str:
    return _time_srt(seconds).replace(',', '.')


def _time_label(seconds: float) -> str:
    s = max(0, int(seconds))
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    return f'{h:02d}:{m:02d}:{sec:02d}' if h else f'{m:02d}:{sec:02d}'


def _line(seg: Segment) -> str:
    who = f'{seg.speaker}：' if seg.speaker else ''
    return f'{who}{seg.text}'


def _write_vtt(result: TranscriptResult, path: Path, note: str | None = None) -> None:
    blocks = ['WEBVTT\n']
    if note:
        blocks.append(f'NOTE {note}\n')
    for s in result.segments:
        end = s.end if s.end > s.start else s.start + 2.0
        blocks.append(f'{_time_vtt(s.start)} --> {_time_vtt(end)}\n{_line(s)}\n')
    path.write_text('\n'.join(blocks), encoding='utf-8-sig')


def _review_html(result: TranscriptResult, title: str, audio_path: str | None, note: str | None) -> str:
    audio_uri = ''
    if audio_path:
        try:
            audio_uri = Path(audio_path).resolve().as_uri()
        except Exception:
            audio_uri = ''
    segments = []
    for i, s in enumerate(result.segments):
        end = s.end if s.end > s.start else s.start + 2.0
        segments.append({
            'i': i, 'start': float(s.start), 'end': float(end),
            'text': s.text, 'speaker': s.speaker or '',
            'label': _time_label(s.start),
        })
    payload = json.dumps(segments, ensure_ascii=False).replace('</', '<\\/')
    note_html = f'<div class="note">{html.escape(note)}</div>' if note else ''
    cleaned_html = ''
    if result.text and str(result.engine).startswith('混合模式'):
        cleaned_html = '<details open><summary><b>Gemini 整理後全文</b></summary><div class="note" style="white-space:pre-wrap;background:#16324a;border-color:#2563eb">' + html.escape(result.text) + '</div></details>'
    return f'''<!doctype html>
<html lang="zh-Hant-TW">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}｜錄音核對播放器</title>
<style>
body{{font-family:"Microsoft JhengHei","Noto Sans TC",sans-serif;margin:0;background:#111827;color:#e5e7eb}}
header{{position:sticky;top:0;background:#111827;padding:18px 22px;border-bottom:1px solid #374151;z-index:5}}
h1{{font-size:20px;margin:0 0 10px}} .tools{{display:flex;gap:10px;flex-wrap:wrap;align-items:center}}
audio{{width:min(920px,100%);height:40px}} input[type=file]{{max-width:320px}}
main{{max-width:1100px;margin:auto;padding:18px 22px 80px}} .note{{background:#4b3b0b;border:1px solid #8a6d1d;padding:10px;border-radius:8px;margin:10px 0}}
#search{{width:min(420px,80vw);padding:8px 10px;border-radius:7px;border:1px solid #4b5563;background:#1f2937;color:#fff}}
.seg{{display:grid;grid-template-columns:86px 1fr;gap:12px;padding:10px 8px;border-bottom:1px solid #293241;cursor:pointer;border-radius:6px}}
.seg:hover{{background:#1f2937}} .seg.active{{background:#123b5d;outline:1px solid #38bdf8}}
.time{{color:#7dd3fc;font-variant-numeric:tabular-nums}} .speaker{{font-weight:700;color:#fcd34d;margin-right:6px}} .text{{line-height:1.7}}
.small{{font-size:12px;color:#9ca3af}} button{{padding:7px 10px;border-radius:7px;border:1px solid #4b5563;background:#1f2937;color:#fff;cursor:pointer}}
</style>
</head>
<body>
<header>
<h1>{html.escape(title)}｜錄音與逐字稿核對</h1>
<div class="tools">
<audio id="audio" controls preload="metadata" src="{html.escape(audio_uri, quote=True)}"></audio>
<label class="small">音檔無法載入時：<input id="pick" type="file" accept="audio/*,video/*"></label>
<input id="search" type="search" placeholder="搜尋逐字稿文字…">
<button id="follow">自動跟隨：開</button>
</div>
<div class="small">點選任一逐字稿段落即可跳到該時間播放。此 HTML 可直接在瀏覽器開啟；若搬移檔案後音訊路徑失效，可用上方「選擇音檔」重新指定。</div>
</header>
<main>{note_html}{cleaned_html}<div id="list"></div></main>
<script>
const segs={payload}; const list=document.getElementById('list'); const audio=document.getElementById('audio');
let follow=true, active=-1;
function esc(s){{return String(s).replace(/[&<>"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[c]));}}
function render(filter=''){{
  const q=filter.trim().toLowerCase(); list.innerHTML='';
  segs.forEach((s,idx)=>{{ if(q && !(s.text+' '+s.speaker).toLowerCase().includes(q)) return;
    const d=document.createElement('div'); d.className='seg'; d.dataset.idx=idx;
    d.innerHTML=`<div class="time">${{esc(s.label)}}</div><div class="text">${{s.speaker?'<span class="speaker">'+esc(s.speaker)+'：</span>':''}}${{esc(s.text)}}</div>`;
    d.onclick=()=>{{audio.currentTime=s.start; audio.play().catch(()=>{{}}); setActive(idx,true);}}; list.appendChild(d);
  }});
}}
function setActive(idx,scroll=false){{
  if(idx===active && !scroll)return; active=idx;
  document.querySelectorAll('.seg.active').forEach(x=>x.classList.remove('active'));
  const el=document.querySelector(`.seg[data-idx="${{idx}}"]`); if(el){{el.classList.add('active'); if(scroll&&follow)el.scrollIntoView({{block:'center',behavior:'smooth'}});}}
}}
audio.addEventListener('timeupdate',()=>{{ const t=audio.currentTime; let idx=-1; for(let i=0;i<segs.length;i++){{if(t>=segs[i].start && t<segs[i].end){{idx=i;break;}}}} if(idx>=0)setActive(idx,true); }});
document.getElementById('search').addEventListener('input',e=>render(e.target.value));
document.getElementById('pick').addEventListener('change',e=>{{const f=e.target.files[0]; if(f){{audio.src=URL.createObjectURL(f);audio.load();}}}});
document.getElementById('follow').onclick=e=>{{follow=!follow;e.target.textContent='自動跟隨：'+(follow?'開':'關');}};
render();
</script>
</body></html>'''


def export_all(result: TranscriptResult, base_path: str, formats: list[str], note: str | None = None,
               source_audio: str | None = None) -> list[str]:
    """Export transcript formats plus an optional clickable audio review player."""
    base = Path(base_path)
    base.parent.mkdir(parents=True, exist_ok=True)
    out = []
    formats = list(dict.fromkeys(formats or []))
    if 'txt' in formats:
        p = base.with_suffix('.txt')
        body = result.text or '\n'.join(_line(x) for x in result.segments)
        if note:
            body = note + '\n\n' + body
        p.write_text(body, encoding='utf-8-sig')
        out.append(str(p))
    if 'srt' in formats and result.segments:
        p = base.with_suffix('.srt')
        blocks = []
        for i, s in enumerate(result.segments, 1):
            end = s.end if s.end > s.start else s.start + 2.0
            blocks.append(f'{i}\n{_time_srt(s.start)} --> {_time_srt(end)}\n{_line(s)}')
        p.write_text('\n\n'.join(blocks), encoding='utf-8-sig')
        out.append(str(p))
    if 'vtt' in formats and result.segments:
        p = base.with_suffix('.vtt')
        _write_vtt(result, p, note)
        out.append(str(p))
    if 'docx' in formats:
        p = base.with_suffix('.docx')
        doc = Document()
        doc.add_heading(base.stem, level=1)
        doc.add_paragraph(f'辨識引擎：{result.engine}')
        if result.language:
            doc.add_paragraph(f'語言：{result.language}')
        if note:
            doc.add_paragraph(note)
        if str(result.engine).startswith('混合模式') and result.text:
            doc.add_heading('Gemini 整理後全文', level=2)
            doc.add_paragraph(result.text)
            if result.segments:
                doc.add_heading('時間軸逐字稿（錄音核對用）', level=2)
        else:
            doc.add_heading('逐字稿', level=2)
        if result.segments:
            for s in result.segments:
                t = f'[{_time_vtt(s.start)[:-4]}] '
                if s.speaker:
                    t += f'{s.speaker}：'
                doc.add_paragraph(t + s.text)
        elif not (str(result.engine).startswith('混合模式') and result.text):
            doc.add_paragraph(result.text)
        doc.save(p)
        out.append(str(p))
    if 'review' in formats and result.segments:
        # Review mode always exports a WebVTT subtitle alongside the HTML player.
        vtt = base.with_name(base.stem + '_字幕').with_suffix('.vtt')
        _write_vtt(result, vtt, note)
        if str(vtt) not in out:
            out.append(str(vtt))
        page = base.with_name(base.stem + '_錄音核對').with_suffix('.html')
        page.write_text(_review_html(result, base.stem, source_audio, note), encoding='utf-8')
        out.append(str(page))
    return out
